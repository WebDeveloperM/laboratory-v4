import os
import re
from io import BytesIO

from docx import Document
from django.conf import settings

PASSPORT_TEMPLATE = os.path.join(
    str(settings.BASE_DIR), 'documents', 'Бензин АИ 92 САУ+ МТБЭ присада паспорт.docx'
)
SPRAVKA_TEMPLATE = os.path.join(
    str(settings.BASE_DIR), 'documents', 'Бензин АИ-91 OZDST 3031нги 01.01.2016 справка.docx'
)


def _set_para_text(para, new_text: str):
    """Collapse all runs into run[0] to preserve paragraph formatting."""
    if not para.runs:
        para.add_run(new_text)
        return
    para.runs[0].text = new_text
    for run in para.runs[1:]:
        run.text = ''


def _apply_replacements(para, replacements):
    """Apply list of (regex_pattern, replacement_str) to paragraph text."""
    text = para.text
    if not text:
        return
    new_text = text
    for pattern, repl in replacements:
        new_text = re.sub(pattern, repl, new_text)
    if new_text != text:
        _set_para_text(para, new_text)


def _fill_table_last_col(table, values: dict, start_row: int = 1):
    """Fill the last column of a table with values dict {str(i): value}."""
    last_col = len(table.columns) - 1
    for i, row in enumerate(table.rows[start_row:]):
        val = values.get(str(i), '')
        if val and len(row.cells) > last_col:
            cell = row.cells[last_col]
            if cell.paragraphs:
                _set_para_text(cell.paragraphs[0], str(val))


def generate_passport_document(passport) -> BytesIO:
    doc = Document(PASSPORT_TEMPLATE)

    pn = passport.passport_number or ''
    reservoir = passport.reservoir or ''
    measurement = passport.measurement or ''
    manufacture_date = passport.manufacture_date or ''
    scd = passport.sample_collection_date or ''
    sad = passport.sample_arrival_date or ''
    test_date = passport.test_date or ''
    batch_number = passport.batch_number or ''
    issue_date = passport.issue_date or ''
    actual = passport.actual_values or {}

    replacements = []

    if pn:
        replacements.append((r'ПАСПОРТ\s*№\s*_+', f'ПАСПОРТ № {pn}'))
    if reservoir:
        replacements.append((r'Резервуар\s*:\s*_+', f'Резервуар: {reservoir}'))
        replacements.append((r'Из резервуара\s*№\s*_+', f'Из резервуара № {reservoir}'))
    if measurement:
        replacements.append((r'Замер\s*:\s*_+', f'Замер: {measurement}'))
    if manufacture_date:
        replacements.append((r'Дата изготовления\s+_+', f'Дата изготовления {manufacture_date}'))
    if scd:
        replacements.append((r'Дата\s+отбора\s+образцов:\s*_+', f'Дата отбора образцов: {scd}'))
    if sad:
        replacements.append((r'Дата поступления\s+образцов:\s*_+', f'Дата поступления образцов: {sad}'))
    if test_date:
        replacements.append((r'Дата проведения\s+испытаний:\s*_+', f'Дата проведения испытаний: {test_date}'))
    if batch_number:
        replacements.append((r'Партия\s*№\s*:\s*_+', f'Партия №: {batch_number}'))
    if issue_date:
        replacements.append((r'Дата выдачи паспорта\s*_+', f'Дата выдачи паспорта {issue_date}'))

    for para in doc.paragraphs:
        _apply_replacements(para, replacements)

    for table in doc.tables:
        _fill_table_last_col(table, actual)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def generate_spravka_document(spravka) -> BytesIO:
    doc = Document(SPRAVKA_TEMPLATE)

    sn = spravka.spravka_number or ''
    manufacture_date = spravka.manufacture_date or ''
    reservoir = spravka.reservoir or ''
    measurement = spravka.measurement or ''
    wagon_count = spravka.wagon_count or ''
    wagon_numbers = spravka.wagon_numbers or ''
    scd = spravka.sample_collection_date or ''
    issue_date = spravka.issue_date or ''
    actual = spravka.actual_values or {}

    replacements = []

    if sn:
        replacements.append((r'Справка\s+№\s*_+', f'Справка № {sn}'))
    if manufacture_date:
        replacements.append((r'Изготовлено\s+_+', f'Изготовлено {manufacture_date}'))
        replacements.append((r'Дата налива\s+в/ц\s*_+', f'Дата налива в/ц {manufacture_date}'))
    if reservoir:
        # Copy 1: "Резервуар ____" (space, no colon)
        replacements.append((r'Резервуар\s+_+', f'Резервуар {reservoir}'))
        # Copies 2-4: "из резервуара №_____"
        replacements.append((r'из резервуара\s*№\s*_+', f'из резервуара № {reservoir}'))
    if measurement:
        replacements.append((r'Замер\s+_+', f'Замер {measurement}'))
    if wagon_count:
        replacements.append((
            r'количество заявленных вагон цистерн\s*_+',
            f'количество заявленных вагон цистерн {wagon_count}'
        ))
    if wagon_numbers:
        replacements.append((r'номера вагон цистерн:\s*_+', f'номера вагон цистерн: {wagon_numbers}'))
    if scd:
        replacements.append((r'Дата проведение отбора\s*_+', f'Дата проведение отбора {scd}'))
    if issue_date:
        replacements.append((r'выдачи паспорта\s*_+', f'выдачи паспорта {issue_date}'))
        replacements.append((r'выдачи справка\s*_+', f'выдачи справка {issue_date}'))

    for para in doc.paragraphs:
        _apply_replacements(para, replacements)

    # Tables 0 and 1: 22 data rows
    for table in doc.tables[:2]:
        _fill_table_last_col(table, actual)

    # Tables 2 and 3: 11 data rows — use first 11 values only
    short_actual = {str(i): actual.get(str(i), '') for i in range(11)}
    for table in doc.tables[2:]:
        _fill_table_last_col(table, short_actual)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
